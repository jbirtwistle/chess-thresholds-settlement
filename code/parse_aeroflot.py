"""Parse the Aeroflot Open 2026 round-9 pairing/results table (chess-results
export, Russian interface) and sanity-check it against the start list.

Run:  python3 -I parse_aeroflot.py
Out:  data/aeroflot2026_r9.csv   (same schema as menorca2026_r9.csv)
"""
import csv
from collections import Counter
from pathlib import Path

import openpyxl

U = Path("/root/.claude/uploads/a5d1dd3e-2c53-5abf-a634-38eb111f5119")
PAIR = U / "1e6cb35c-chessResultsList_1.xlsx"
START = U / "968ad1da-chessResultsList.xlsx"
OUT = Path("/home/claude/chess_sim/data/aeroflot2026_r9.csv")


def pts(x):
    """6 -> 6.0 ; '5½' -> 5.5 ; '½' -> 0.5."""
    if x in (None, ""):
        return None
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip()
    half = "½" in s
    s = s.replace("½", "").strip()
    return (float(s) if s else 0.0) + (0.5 if half else 0.0)


def kind_and_points(res, black):
    r = str(res).strip()
    if black in ("bye", "без пары"):
        return ("bye", 1.0, None) if black == "bye" and r == "1" else ("absent", 0.0, None)
    table = {"½ - ½": (0.5, 0.5, "played"), "1 - 0": (1.0, 0.0, "played"),
             "0 - 1": (0.0, 1.0, "played"), "+ - -": (1.0, 0.0, "forfeit"),
             "- - +": (0.0, 1.0, "forfeit")}
    if r not in table:
        raise SystemExit(f"unrecognised result {r!r}")
    w, b, k = table[r]
    return k, w, b


def main():
    rows = list(openpyxl.load_workbook(PAIR, data_only=True).active.iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows) if r[0] == "Bo.")
    print("round header:", rows[hdr - 1][0])
    out = []
    for r in rows[hdr + 1:]:
        if not isinstance(r[0], int):
            continue
        black_name = r[10]
        kind, wp, bp = kind_and_points(r[7], black_name)
        has_opp = kind in ("played", "forfeit")
        res = str(r[7]).strip().replace(" ", "").replace("½-½", "½-½")
        out.append({
            "board": r[0],
            "white": r[4], "white_title": r[3] or "", "white_rank": r[1],
            "white_score_before": pts(r[6]), "white_elo": int(r[5]) if r[5] else "", "white_fed": "",
            "black": black_name if has_opp else "", "black_title": (r[9] or "") if has_opp else "",
            "black_rank": r[13] if has_opp else "",
            "black_score_before": pts(r[8]) if has_opp else "",
            "black_elo": int(r[11]) if has_opp and r[11] else "", "black_fed": "",
            "result": res, "kind": kind,
            "white_pts": wp, "black_pts": bp if bp is not None else "",
            "draw": int(kind == "played" and wp == 0.5),
        })
    print(f"{len(out)} rows; by kind: {dict(Counter(o['kind'] for o in out))}")

    # --- sanity checks -----------------------------------------------------
    played = [o for o in out if o["kind"] == "played"]
    diffs = Counter(abs(o["white_score_before"] - o["black_score_before"]) for o in played)
    print("score difference between paired players (Swiss => mostly 0 or 0.5):",
          dict(sorted(diffs.items())))
    ids = [o["white_rank"] for o in out] + [o["black_rank"] for o in out if o["black_rank"] != ""]
    print(f"each start number at most once: {len(ids) == len(set(ids))} ({len(ids)} players listed)")

    sl = list(openpyxl.load_workbook(START, data_only=True).active.iter_rows(values_only=True))
    sh = next(i for i, r in enumerate(sl) if r[0] == "Ном.")
    start = {r[0]: (r[3], r[6]) for r in sl[sh + 1:] if isinstance(r[0], int)}
    mism = n = 0
    for o in out:
        for side in ("white", "black"):
            rk = o[f"{side}_rank"]
            if rk in start and o[side]:
                n += 1
                if start[rk][0].strip() != o[side].strip() or (
                        o[f"{side}_elo"] != "" and int(start[rk][1]) != int(o[f"{side}_elo"])):
                    mism += 1
                    print("   MISMATCH vs start list:", rk, o[side], start[rk])
    print(f"start-list cross-check: {n} players compared, {mism} mismatches "
          f"(start list has {len(start)} entries, pairing refers to start numbers up to "
          f"{max(i for i in ids)})")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(out[0]))
        w.writeheader()
        w.writerows(out)
    print("wrote", OUT)


if __name__ == "__main__":
    main()

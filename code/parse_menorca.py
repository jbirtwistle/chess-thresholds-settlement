"""Parse the Menorca 2026 round-9 pairing table and cross-check it against the
final standings table (both user-supplied exports from info64.org).

Run:  python3 -I parse_menorca.py
Out:  data/menorca2026_r9.csv
"""
import csv
import re
from pathlib import Path

import openpyxl

U = Path("/root/.claude/uploads/a5d1dd3e-2c53-5abf-a634-38eb111f5119")
PAIR = U / "86e1c8ea-2026-10-06-18-12-1791303156-pairing-table.xlsx"
STAND = U / "84a8b2b4-2026-10-06-17-59-1791302384-standings-table.xlsx"
OUT = Path("/home/claude/chess_sim/data/menorca2026_r9.csv")


def num(x):
    """'7,5' -> 7.5 ; '2595' -> 2595.0 ; None/'' -> None."""
    if x in (None, ""):
        return None
    return float(str(x).replace(",", ".").strip())


def strip_title(name: str):
    m = re.match(r"^(GM|IM|FM|WGM|WIM|WFM|CM|WCM)\s+(.*)$", name.strip())
    return (m.group(1), m.group(2)) if m else ("", name.strip())


def result_points(res: str, has_opponent: bool):
    """Return (white_pts, black_pts, kind).

    kind: 'played' (a game over the board), 'forfeit' (one side did not show),
    'bye' (point without an opponent) or 'absent' (no opponent, 0 points).
    Forfeit/bye/absent conventions are checked against the final standings.
    """
    r = res.strip()
    if has_opponent:
        if r in ("½-½", "1/2-1/2"):
            return 0.5, 0.5, "played"
        if r == "1-0":
            return 1.0, 0.0, "played"
        if r == "0-1":
            return 0.0, 1.0, "played"
        if r == "+--":
            return 1.0, 0.0, "forfeit"
        if r == "--+":
            return 0.0, 1.0, "forfeit"
        if r == "---":
            return 0.0, 0.0, "forfeit"
        raise SystemExit(f"unrecognised result with opponent: {r!r}")
    if r == "+":
        return 1.0, None, "bye"
    if r == "0":
        return 0.0, None, "absent"
    raise SystemExit(f"unrecognised result without opponent: {r!r}")


def main():
    ws = openpyxl.load_workbook(PAIR, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r[0] == "M.")
    boards, odd = [], []
    for r in rows[hdr_i + 1:]:
        if r[0] in (None, "") or not str(r[0]).strip().isdigit():
            if any(x not in (None, "") for x in r):
                odd.append(r)           # byes / footers
            continue
        has_opp = r[7] not in (None, "")
        wt, wn = strip_title(r[1])
        bt, bn = strip_title(r[7]) if has_opp else ("", "")
        wp, bp, kind = result_points(str(r[6]), has_opp)
        res = str(r[6]).strip()
        boards.append({
            "board": int(r[0]),
            "white": wn, "white_title": wt, "white_rank": int(r[2]),
            "white_score_before": num(r[3]), "white_elo": num(r[4]), "white_fed": r[5],
            "black": bn, "black_title": bt,
            "black_rank": int(r[8]) if has_opp else "",
            "black_score_before": num(r[9]) if has_opp else "",
            "black_elo": num(r[10]) if has_opp else "", "black_fed": r[11] if has_opp else "",
            "result": res, "kind": kind,
            "white_pts": wp, "black_pts": bp if bp is not None else "",
            "draw": int(res in ("½-½", "1/2-1/2")) if kind == "played" else "",
        })
    from collections import Counter
    print(f"{len(boards)} rows parsed; by kind: {dict(Counter(b['kind'] for b in boards))}; "
          f"non-row lines: {len(odd)}")

    # ---- cross-check with the final standings -----------------------------
    ws2 = openpyxl.load_workbook(STAND, data_only=True).active
    r2 = list(ws2.iter_rows(values_only=True))
    h2 = next(i for i, r in enumerate(r2) if r[0] == "Pos.")
    final = {}
    for r in r2[h2 + 1:]:
        if r[1] not in (None, "") and str(r[1]).strip().isdigit() and r[5] not in (None, ""):
            final[int(r[1])] = num(r[5])
    bad, checked, by_kind = [], 0, Counter()
    for b in boards:
        for side in ("white", "black"):
            pts = b[f"{side}_pts"]
            if pts == "" or (side == "black" and b["black_rank"] == ""):
                continue
            rk = b[f"{side}_rank"]
            exp = b[f"{side}_score_before"] + pts
            checked += 1
            if rk not in final:
                by_kind["not in standings"] += 1
            elif abs(final[rk] - exp) > 1e-9:
                bad.append((b[side], rk, b["kind"], exp, final[rk]))
                by_kind[b["kind"]] += 1
    print(f"cross-check vs final standings: {checked} player-results checked, "
          f"{len(bad)} mismatches {dict(by_kind)}")
    for x in bad[:12]:
        print("   MISMATCH (name, rank, kind, before+result, final):", x)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(boards[0]))
        w.writeheader()
        w.writerows(boards)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
